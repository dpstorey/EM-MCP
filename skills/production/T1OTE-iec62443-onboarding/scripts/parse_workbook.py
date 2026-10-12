#!/usr/bin/env python3
# SPDX-License-Identifier: Apache-2.0
"""Load a filled onboarding workbook (xlsx), a CSV directory, or a JSON model
into the normalised model.

usage: parse_workbook.py INPUT [--out model.json]
  INPUT: .xlsx | directory with site.csv zones.csv address_space.csv conduits.csv options.csv | .json

Prints {"model": {...}, "gaps": [...]} . Gaps are questions to ask the user
(blank required cells); structural validation is validate_model.py's job.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))
from iec_common import dump_json, load_json, normalize_model  # noqa: E402

SHEET_KEYS = {
    "site": ("site",),
    "zones": ("zones",),
    "address_space": ("address space", "address_space", "addressspace"),
    "conduits": ("conduits",),
    "options": ("options",),
}
FIRST_HEADER = {
    "site": "site name",
    "zones": "zone name",
    "address_space": "zone",
    "conduits": "conduit name",
    "options": "option",
}


def _key(h: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(h or "").strip().lower()).strip("_")


def _rows_from_grid(grid: list[list[Any]], first_header: str) -> list[dict[str, Any]]:
    hdr_idx = next(
        (i for i, r in enumerate(grid) if r and str(r[0] or "").strip().lower() == first_header), None
    )
    if hdr_idx is None:
        return []
    headers = [_key(h) for h in grid[hdr_idx]]
    out = []
    for r in grid[hdr_idx + 1 :]:
        if not any(str(c or "").strip() for c in r):
            continue
        if str(r[0] or "").strip().lower().startswith("(example)"):
            continue
        out.append({h: (r[i] if i < len(r) else None) for i, h in enumerate(headers) if h})
    return out


def _read_xlsx(path: Path) -> dict[str, list[dict[str, Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    by_lower = {ws.title.strip().lower(): ws for ws in wb.worksheets}
    out: dict[str, list[dict[str, Any]]] = {}
    for key, names in SHEET_KEYS.items():
        ws = next((by_lower[n] for n in names if n in by_lower), None)
        out[key] = (
            _rows_from_grid([list(r) for r in ws.iter_rows(values_only=True)], FIRST_HEADER[key]) if ws else []
        )
    return out


def _read_csv_dir(path: Path) -> dict[str, list[dict[str, Any]]]:
    out: dict[str, list[dict[str, Any]]] = {}
    for key in SHEET_KEYS:
        f = path / f"{key}.csv"
        if not f.exists():
            out[key] = []
            continue
        with open(f, newline="", encoding="utf-8-sig") as fh:
            out[key] = _rows_from_grid(list(csv.reader(fh)), FIRST_HEADER[key])
    return out


OPTION_LABELS = {
    "create_sample_groups": "create_sample_groups",
    "create_sample_policies": "create_sample_policies",
    "policies_enabled": "policies_enabled",
    "severity": "severity",
    "name_prefix": "name_prefix",
    "direction_tolerant": "direction_tolerant",
    "allow_zone_overlap": "allow_overlap",
}


def tables_to_model(t: dict[str, list[dict[str, Any]]]) -> tuple[dict[str, Any], list[str]]:
    gaps: list[str] = []
    site_row = (t["site"] or [{}])[0]
    raw = {
        "site": {"name": site_row.get("site_name"), "customer": site_row.get("customer"), "notes": site_row.get("notes")},
        "zones": [
            {
                "name": r.get("zone_name"),
                "description": r.get("description"),
                "purdue_level": r.get("purdue_level"),
                "sl_target": r.get("sl_t"),
                "criticality": r.get("criticality"),
            }
            for r in t["zones"]
        ],
        "members": [
            {"zone": r.get("zone"), "method": r.get("method"), "value": r.get("value"), "end": r.get("range_end"), "comment": r.get("comment")}
            for r in t["address_space"]
        ],
        "conduits": [
            {
                "name": r.get("conduit_name"),
                "zone_a": r.get("zone_a"),
                "zone_b": r.get("zone_b"),
                "direction": r.get("direction"),
                "allowed_protocols": r.get("allowed_protocols"),
                "description": r.get("comment"),
            }
            for r in t["conduits"]
        ],
        "options": {},
    }
    for r in t["options"]:
        k = OPTION_LABELS.get(_key(r.get("option")))
        if k:
            raw["options"][k] = r.get("value")
    model = normalize_model(raw)
    if not model["site"]["name"]:
        gaps.append("Site name is blank - which site (as paired in Enterprise Manager) is this for?")
    if not model["zones"]:
        gaps.append("No zones found on the Zones sheet.")
    for i, m in enumerate(model["members"], start=1):
        if not m["zone"] or not m["method"] or not m["value"]:
            gaps.append(f"Address Space row {i}: zone, method and value are all required.")
    for c in model["conduits"]:
        who = c["name"] or f"{c['zone_a']}/{c['zone_b']}"
        if not c["direction"]:
            gaps.append(f"Conduit '{who}': direction is blank - which zone may initiate (a_to_b, b_to_a, or bidirectional)?")
        if not c["allowed_protocols"]:
            gaps.append(f"Conduit '{who}': no allowed protocols - no conformance check can be generated.")
    return model, gaps


def parse(path: str) -> dict[str, Any]:
    p = Path(path)
    if p.is_dir():
        model, gaps = tables_to_model(_read_csv_dir(p))
    elif p.suffix.lower() == ".json":
        model, gaps = normalize_model(load_json(p)), []
    elif p.suffix.lower() in (".xlsx", ".xlsm"):
        model, gaps = tables_to_model(_read_xlsx(p))
    else:
        raise SystemExit(f"unsupported input {path!r}: use .xlsx, .json or a CSV directory")
    return {"model": model, "gaps": gaps}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("input")
    ap.add_argument("--out")
    a = ap.parse_args()
    dump_json(parse(a.input), a.out)


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
# SPDX-License-Identifier: Apache-2.0
"""Generate the onboarding workbook (xlsx) with dropdowns and example rows.

usage: generate_template.py [--site "Site name"] [--out onboarding-template.xlsx]

Rows whose first cell starts with "(example)" are ignored by parse_workbook.py,
so the examples can stay in the file. Requires openpyxl.
"""

from __future__ import annotations

import argparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HEADER_ROW = 5
FIRST_DATA_ROW = 6
LAST_DATA_ROW = 205
FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
EX_FILL = PatternFill("solid", fgColor="EDEDED")

SHEETS: dict[str, dict] = {
    "Site": {
        "title": "Site",
        "help": "One row. The site name must match a paired site in Enterprise Manager; the agent confirms it.",
        "cols": ["Site name", "Customer", "Notes"],
        "widths": [30, 30, 60],
        "example": ["(example) Plant 1", "Example Corp", "Pilot site"],
    },
    "Zones": {
        "title": "Zones (IEC 62443 'zone' = a set of assets, i.e. an asset group)",
        "help": "One row per zone. Names must be unique. Purdue 0-5 and SL-T 1-4 are optional.",
        "cols": ["Zone name", "Description", "Purdue level", "SL-T", "Criticality"],
        "widths": [28, 50, 14, 10, 14],
        "example": ["(example) Control", "PLC and I/O network", 1, 2, "high"],
    },
    "Address Space": {
        "title": "Address space per zone (a zone may have several rows)",
        "help": (
            "Method: cidr (e.g. 10.10.1.0/24), range (start in Value, end in Range end), "
            "ip_list (comma separated), vlan (id), asset_names (agent will ask you to pick), filter (JSON)."
        ),
        "cols": ["Zone", "Method", "Value", "Range end", "Comment"],
        "widths": [28, 14, 40, 18, 40],
        "example": ["(example) Control", "cidr", "10.10.1.0/24", "", "Line 1 PLC VLAN"],
    },
    "Conduits": {
        "title": "Conduits (IEC 62443 'conduit' = allowed communication path, i.e. detection policies)",
        "help": (
            "Direction is who may INITIATE: a_to_b (A starts, B answers), b_to_a, or bidirectional "
            "(only when both sides really initiate). Direction is enforced; choose carefully. "
            "Allowed protocols: comma separated (Modbus, S7, CIP, EtherNet/IP, DNP3, IEC 104, IEC 61850, "
            "PROFINET, OPC UA, HTTPS, HTTP, SSH, RDP, NTP, DNS, SNMP, Syslog, BACnet)."
        ),
        "cols": ["Conduit name", "Zone A", "Zone B", "Direction", "Allowed protocols", "Comment"],
        "widths": [28, 24, 24, 16, 44, 40],
        "example": ["(example) DMZ to Control", "DMZ", "Control", "a_to_b", "OPC UA, HTTPS", "Historian pulls from PLCs"],
    },
    "Options": {
        "title": "Options",
        "help": "Safe defaults shown. Policies are created disabled unless you say otherwise.",
        "cols": ["Option", "Value", "Meaning"],
        "widths": [30, 18, 80],
        "rows": [
            ["Create sample groups", "no", "Sample asset groups by type (PLCs, HMIs, ...) for browsing"],
            ["Create sample policies", "no", "Remote-admin-to-controllers and new-asset-in-zone policies"],
            ["Policies enabled", "no", "yes = create enabled; no = create then disable (recommended for first run)"],
            ["Severity", "medium", "low | medium | high - level of generated policies"],
            ["Name prefix", "IEC62443", "Prefix on every group and policy the agent creates"],
            ["Direction tolerant", "no", "yes = treat every conduit as symmetric for boundary checks (fewer false positives, weaker)"],
            ["Allow zone overlap", "no", "yes = permit overlapping address space between zones (not recommended)"],
        ],
    },
}


def _style_sheet(ws, spec: dict) -> None:
    ws["A1"] = spec["title"]
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = spec["help"]
    ws["A2"].font = Font(name=FONT, italic=True)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=max(len(spec["cols"]), 4))
    ws.row_dimensions[2].height = 30
    ws["A4"] = "Yellow cells are yours to edit. Grey '(example)' rows are ignored; keep or delete them."
    ws["A4"].font = Font(name=FONT, size=9)
    for i, (name, w) in enumerate(zip(spec["cols"], spec["widths"]), start=1):
        c = ws.cell(row=HEADER_ROW, column=i, value=name)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=1)


def build(site: str | None, out: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    readme = wb.create_sheet("README")
    lines = [
        ("IEC 62443 onboarding workbook for Tenable OT", True),
        ("Zones are modelled as asset groups (horizontal: which assets). Conduits are modelled as detection policies (vertical: which paths are allowed).", False),
        ("Fill in: Site, Zones, Address Space, Conduits, Options. Then give the file to the onboarding agent.", False),
        ("The agent validates it, shows a plan, and changes nothing until you approve.", False),
        ("Direction is strict: a conduit a_to_b means zone A may initiate to zone B and not the reverse.", False),
        ("Yellow = input. Grey '(example)' rows are ignored.", False),
    ]
    for r, (t, b) in enumerate(lines, start=1):
        readme.cell(row=r, column=1, value=t).font = Font(name=FONT, bold=b, size=14 if b else 11)
    readme.column_dimensions["A"].width = 140

    sheets = {}
    for name, spec in SHEETS.items():
        ws = wb.create_sheet(name)
        sheets[name] = ws
        _style_sheet(ws, spec)
        if "rows" in spec:
            for r, row in enumerate(spec["rows"], start=FIRST_DATA_ROW):
                for c, v in enumerate(row, start=1):
                    cell = ws.cell(row=r, column=c, value=v)
                    cell.font = Font(name=FONT)
                    if c == 2:
                        cell.fill = INPUT_FILL
            continue
        for c, v in enumerate(spec["example"], start=1):
            cell = ws.cell(row=FIRST_DATA_ROW, column=c, value=v)
            cell.font = Font(name=FONT, italic=True, color="595959")
            cell.fill = EX_FILL
        last = FIRST_DATA_ROW if name == "Site" else LAST_DATA_ROW
        for r in range(FIRST_DATA_ROW + 1, last + 1):
            for c in range(1, len(spec["cols"]) + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = INPUT_FILL
                cell.font = Font(name=FONT)
    if site:
        ws = sheets["Site"]
        ws.cell(row=FIRST_DATA_ROW + 1, column=1, value=site).font = Font(name=FONT)
        ws.cell(row=FIRST_DATA_ROW + 1, column=1).fill = INPUT_FILL

    def dv(ws, formula: str, col: str, rows=(FIRST_DATA_ROW + 1, LAST_DATA_ROW), blank=True):
        v = DataValidation(type="list", formula1=formula, allow_blank=blank, showErrorMessage=True)
        ws.add_data_validation(v)
        v.add(f"{col}{rows[0]}:{col}{rows[1]}")

    zone_rng = f"=Zones!$A${FIRST_DATA_ROW}:$A${LAST_DATA_ROW}"
    dv(sheets["Address Space"], '"cidr,range,ip_list,vlan,asset_names,filter"', "B")
    dv(sheets["Address Space"], zone_rng, "A")
    dv(sheets["Conduits"], zone_rng, "B")
    dv(sheets["Conduits"], zone_rng, "C")
    dv(sheets["Conduits"], '"a_to_b,b_to_a,bidirectional"', "D")
    dv(sheets["Zones"], '"0,1,2,3,4,5"', "C")
    dv(sheets["Zones"], '"1,2,3,4"', "D")
    dv(sheets["Zones"], '"low,medium,high"', "E")
    opt = sheets["Options"]
    yn = DataValidation(type="list", formula1='"yes,no"', allow_blank=False)
    opt.add_data_validation(yn)
    for r in (FIRST_DATA_ROW, FIRST_DATA_ROW + 1, FIRST_DATA_ROW + 2, FIRST_DATA_ROW + 5, FIRST_DATA_ROW + 6):
        yn.add(f"B{r}")
    sev = DataValidation(type="list", formula1='"low,medium,high"', allow_blank=False)
    opt.add_data_validation(sev)
    sev.add(f"B{FIRST_DATA_ROW + 3}")
    wb.save(out)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--site", help="pre-fill the site name")
    ap.add_argument("--out", default="onboarding-template.xlsx")
    a = ap.parse_args()
    build(a.site, a.out)
    print(f"wrote {a.out}")


if __name__ == "__main__":
    main()
